import io
from typing import List
from docx import Document as DocxDocument
from app.core.exceptions import DocMindException
from app.ingestion.normalizer import normalize_text
from app.ingestion.parsers.base import ExtractedBlock, ParsedDocument
from app.models.enums import DocumentErrorCode


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as e:
        raise DocMindException(
            status_code=400,
            code=DocumentErrorCode.CORRUPT_FILE.value,
            title="Corrupt DOCX Document",
            detail=f"Failed to read DOCX file: {str(e)}"
        )

    blocks: List[ExtractedBlock] = []
    current_section = "Introduction"
    total_chars = 0

    # Extract paragraphs and detect heading hierarchy
    for p in doc.paragraphs:
        raw_text = p.text
        if not raw_text.strip():
            continue

        style_name = (p.style.name.lower() if p.style and p.style.name else "")
        is_heading = "heading" in style_name or "title" in style_name
        cleaned = normalize_text(raw_text)

        if is_heading:
            current_section = cleaned[:100]

        block = ExtractedBlock(
            text=cleaned,
            section_path=current_section,
            is_heading=is_heading,
            char_count=len(cleaned)
        )
        blocks.append(block)
        total_chars += len(cleaned)

    # Extract tables in document
    for table_idx, table in enumerate(doc.tables, 1):
        table_rows = []
        for row in table.rows:
            row_text = [normalize_text(cell.text) for cell in row.cells]
            table_rows.append(" | ".join(row_text))

        table_content = "\n".join(table_rows)
        if table_content.strip():
            block = ExtractedBlock(
                text=table_content,
                section_path=f"{current_section} > Table {table_idx}",
                is_heading=False,
                char_count=len(table_content)
            )
            blocks.append(block)
            total_chars += len(table_content)

    if total_chars < 30:
        raise DocMindException(
            status_code=422,
            code=DocumentErrorCode.NO_EXTRACTABLE_TEXT.value,
            title="No Extractable Text",
            detail="The DOCX document is empty or contains no readable text."
        )

    return ParsedDocument(
        blocks=blocks,
        page_count=max(1, len(blocks) // 5),
        total_char_count=total_chars,
        parser_name="python-docx",
        detected_language="en"
    )
